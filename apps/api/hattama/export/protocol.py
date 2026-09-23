"""Protocol document model (one consistent version for registry, summary and exports) and PDF/DOCX
renderers. Fonts: bundled DejaVu Sans (covers Ә Ғ Қ Ң Ө Ұ Ү Һ І); no remote fonts."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from hattama.db.models import ActionItem, Meeting, Participant, ReviewIssue, SummaryRevision

FONT_DIR = Path(__file__).resolve().parent / "fonts"
EXEC_LABEL = {"open": "открыто", "in_progress": "в работе", "done": "выполнено", "cancelled": "отменено"}
REVIEW_LABEL = {"proposed": "предложено", "needs_review": "требует проверки", "confirmed": "подтверждено",
                "edited": "исправлено", "rejected": "отклонено"}


def deadline_text(d: dict[str, Any] | None) -> str:
    if not d or d.get("kind") in (None, "missing"):
        return "срок не указан"
    raw = d.get("raw_text") or ""
    if d.get("kind") == "event":
        return f"после события: {d.get('anchor_event') or raw}"
    if d.get("normalized_date"):
        mark = " (?)" if d.get("ambiguous") or d.get("conflict") else ""
        return f"{d['normalized_date']}{mark} — «{raw}»"
    return f"«{raw}» (дата не определена)"


def person(v: dict[str, Any] | None) -> str:
    return (v or {}).get("name") or "—"


def build_snapshot(session: Session, meeting: Meeting) -> dict[str, Any]:
    participants = session.scalars(select(Participant).where(Participant.meeting_id == meeting.id)).all()
    actions = session.scalars(select(ActionItem).where(ActionItem.meeting_id == meeting.id,
                                                       ActionItem.review_state != "rejected")
                              .order_by(ActionItem.number)).all()
    summary = session.get(SummaryRevision, meeting.active_summary_id) if meeting.active_summary_id else None
    issues = session.scalars(select(ReviewIssue).where(ReviewIssue.meeting_id == meeting.id,
                                                       ReviewIssue.status == "open")).all()
    content = summary.content if summary else {}
    return {
        "meeting": {"id": meeting.id, "title": meeting.title, "date": meeting.meeting_date.isoformat(),
                    "start_time": meeting.start_time, "timezone": meeting.timezone, "status": meeting.status,
                    "protocol_version": meeting.protocol_version},
        "participants": [{"name": p.display_name, "position": p.position, "present": p.is_present}
                         for p in participants],
        "topics": content.get("topics", []),
        "summary": {k: [{"text": i["text"], "flags": i.get("flags", [])} for i in content.get(k, [])]
                    for k in ("facts", "decisions", "assumptions", "risks", "open_questions")},
        "actions": [{"number": a.number, "action": a.action, "assignee": person(a.assignee),
                     "assignee_type": a.assignee_type, "curator": person(a.curator),
                     "deadline": deadline_text(a.deadline), "condition": a.condition, "conditional": a.conditional,
                     "review_state": a.review_state, "execution_state": a.execution_state, "on_hold": a.on_hold,
                     "version": a.version} for a in actions],
        "open_questions": [i.question for i in issues],
        "transcript_revision_id": meeting.active_transcript_revision_id,
        "summary_revision_id": meeting.active_summary_id,
    }


def snapshot_hash(snapshot: dict[str, Any]) -> str:
    return hashlib.sha256(json.dumps(snapshot, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


SECTIONS = [("facts", "Факты и показатели"), ("decisions", "Принятые решения"),
            ("assumptions", "Предположения и оценки участников"), ("risks", "Риски"),
            ("open_questions", "Открытые вопросы из обсуждения")]


def render_pdf(snap: dict[str, Any], path: Path, draft: bool, approval: str | None) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_DIR / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
    base = ParagraphStyle("b", fontName="DejaVu", fontSize=9.5, leading=12, wordWrap="CJK")
    h1 = ParagraphStyle("h1", parent=base, fontName="DejaVu-Bold", fontSize=15, leading=19)
    h2 = ParagraphStyle("h2", parent=base, fontName="DejaVu-Bold", fontSize=11.5, leading=15, spaceBefore=8)
    cell = ParagraphStyle("c", parent=base, fontSize=8.5, leading=10.5)
    m = snap["meeting"]

    def on_page(canvas, doc):  # type: ignore[no-untyped-def]
        canvas.saveState()
        canvas.setFont("DejaVu", 7.5)
        status = "ЧЕРНОВИК — НЕ УТВЕРЖДЁН" if draft else f"Утверждено: версия {m['protocol_version']} ({approval})"
        canvas.drawString(15 * mm, 8 * mm, f"{m['title']} · {status} · стр. {doc.page}")
        if draft:
            canvas.setFont("DejaVu-Bold", 60)
            canvas.setFillColor(colors.Color(0.85, 0.2, 0.2, alpha=0.12))
            canvas.translate(doc.pagesize[0] / 2, doc.pagesize[1] / 2)
            canvas.rotate(35)
            canvas.drawCentredString(0, 0, "ЧЕРНОВИК")
        canvas.restoreState()

    esc = lambda t: (str(t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))  # noqa: E731
    story: list[Any] = [Paragraph(f"Протокол совещания: {esc(m['title'])}", h1),
                        Paragraph(f"Дата: {m['date']} {m['start_time'] or ''} · часовой пояс {m['timezone']} · "
                                  f"версия протокола {m['protocol_version']} · "
                                  f"{'ЧЕРНОВИК' if draft else 'УТВЕРЖДЁН'}", base), Spacer(1, 6)]
    story.append(Paragraph("Участники", h2))
    story.append(Paragraph("; ".join(f"{esc(p['name'])}{' (' + esc(p['position']) + ')' if p['position'] else ''}"
                                     f"{'' if p['present'] else ' — отсутствовал(а)'}" for p in snap["participants"])
                           or "—", base))
    if snap["topics"]:
        story += [Paragraph("Темы", h2), Paragraph("; ".join(esc(t) for t in snap["topics"]), base)]
    for key, title in SECTIONS:
        items = snap["summary"].get(key) or []
        if items:
            story.append(Paragraph(title, h2))
            for it in items:
                flag = f" <font color='#b45309'>[проверить: {esc('; '.join(it['flags']))}]</font>" if it["flags"] else ""
                story.append(Paragraph(f"• {esc(it['text'])}{flag}", base))
    story.append(Paragraph("Поручения", h2))
    rows = [[Paragraph(h, cell) for h in ("№", "Поручение", "Исполнитель", "Срок", "Условие", "Проверка", "Статус")]]
    for a in snap["actions"]:
        rows.append([Paragraph(str(a["number"]), cell), Paragraph(esc(a["action"]), cell),
                     Paragraph(esc(a["assignee"]) + (f"<br/>куратор: {esc(a['curator'])}" if a["curator"] != "—"
                                                     else ""), cell),
                     Paragraph(esc(a["deadline"]), cell),
                     Paragraph(esc(a["condition"] or ("условное" if a["conditional"] else "—")), cell),
                     Paragraph(REVIEW_LABEL.get(a["review_state"], a["review_state"]), cell),
                     Paragraph(EXEC_LABEL.get(a["execution_state"], a["execution_state"]) +
                               (" (приостановлено)" if a["on_hold"] else ""), cell)])
    table = LongTable(rows, repeatRows=1, colWidths=[10 * mm, 80 * mm, 40 * mm, 45 * mm, 45 * mm, 25 * mm, 25 * mm])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.92, 0.92, 0.95)),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(table if len(rows) > 1 else Paragraph("Поручений нет.", base))
    if snap["open_questions"]:
        story.append(Paragraph("Вопросы на уточнение (не закрыты)", h2))
        for q in snap["open_questions"]:
            story.append(Paragraph(f"• {esc(q)}", base))
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=m["title"])
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


def render_docx(snap: dict[str, Any], path: Path, draft: bool, approval: str | None) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    style = doc.styles["Normal"]
    style.font.name = "DejaVu Sans"
    style.font.size = Pt(10)
    m = snap["meeting"]
    header = sec.header.paragraphs[0]
    header.text = ("ЧЕРНОВИК — НЕ УТВЕРЖДЁН" if draft
                   else f"Утверждено: версия протокола {m['protocol_version']} ({approval})")
    if draft:
        for run in header.runs:
            run.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)
            run.bold = True
    doc.add_heading(f"Протокол совещания: {m['title']}", 1)
    doc.add_paragraph(f"Дата: {m['date']} {m['start_time'] or ''} · часовой пояс {m['timezone']} · "
                      f"версия {m['protocol_version']}")
    doc.add_heading("Участники", 2)
    doc.add_paragraph("; ".join(p["name"] + ("" if p["present"] else " (отсутствовал(а))")
                                for p in snap["participants"]) or "—")
    if snap["topics"]:
        doc.add_heading("Темы", 2)
        doc.add_paragraph("; ".join(snap["topics"]))
    for key, title in SECTIONS:
        items = snap["summary"].get(key) or []
        if items:
            doc.add_heading(title, 2)
            for it in items:
                doc.add_paragraph(it["text"] + (f" [проверить: {'; '.join(it['flags'])}]" if it["flags"] else ""),
                                  style="List Bullet")
    doc.add_heading("Поручения", 2)
    cols = ("№", "Поручение", "Исполнитель", "Срок", "Условие", "Проверка", "Статус")
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = "Table Grid"
    for i, h in enumerate(cols):
        table.rows[0].cells[i].text = h
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for a in snap["actions"]:
        cells = table.add_row().cells
        values = (str(a["number"]), a["action"], a["assignee"] + (f"\nкуратор: {a['curator']}" if a["curator"] != "—"
                                                                   else ""),
                  a["deadline"], a["condition"] or ("условное" if a["conditional"] else "—"),
                  REVIEW_LABEL.get(a["review_state"], a["review_state"]),
                  EXEC_LABEL.get(a["execution_state"], a["execution_state"]))
        for i, v in enumerate(values):
            cells[i].text = v
    if snap["open_questions"]:
        doc.add_heading("Вопросы на уточнение (не закрыты)", 2)
        for q in snap["open_questions"]:
            doc.add_paragraph(q, style="List Bullet")
    doc.save(str(path))
